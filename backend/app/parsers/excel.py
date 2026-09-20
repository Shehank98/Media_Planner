"""Thin openpyxl helpers shared by the parsers.

We read headers PER SHEET (never assume a fixed column layout) and match them
tolerantly (case-insensitive, whitespace-normalised, alias-aware).
"""
from __future__ import annotations

import re
from typing import Any

from openpyxl import load_workbook


def normalise_header(h: Any) -> str:
    if h is None:
        return ""
    return re.sub(r"\s+", " ", str(h).strip()).lower()


def build_header_index(header_row: list[Any]) -> dict[str, int]:
    """Map normalised header -> column index."""
    idx: dict[str, int] = {}
    for i, h in enumerate(header_row):
        key = normalise_header(h)
        if key and key not in idx:
            idx[key] = i
    return idx


def find_column(header_index: dict[str, int], aliases: list[str]) -> int | None:
    """Return the column index for the first matching alias.

    Matches exact normalised header first, then a substring/contains match so
    'rack rate 30 sec' resolves for the alias 'rack rate'."""
    for alias in aliases:
        a = normalise_header(alias)
        if a in header_index:
            return header_index[a]
    for alias in aliases:
        a = normalise_header(alias)
        for header, i in header_index.items():
            if a and a in header:
                return i
    return None


def load_sheets(path: str) -> dict[str, list[list[Any]]]:
    """Return {sheet_name: rows} where rows is a list of cell-value lists.

    Opened as a binary file object rather than by path so openpyxl reads the
    zip directly and does not reject staged uploads that were saved without an
    .xlsx extension.
    """
    out: dict[str, list[list[Any]]] = {}
    with open(path, "rb") as fh:
        wb = load_workbook(fh, data_only=True, read_only=True)
        try:
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                out[ws.title] = rows
        finally:
            wb.close()
    return out


def to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = re.sub(r"[^\d.\-]", "", str(value))
    if s in ("", "-", ".", "-."):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def to_int(value: Any) -> int | None:
    f = to_float(value)
    return int(f) if f is not None else None


def cell(row: list[Any], idx: int | None) -> Any:
    if idx is None or idx >= len(row):
        return None
    return row[idx]
